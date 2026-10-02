"""Reading a spreadsheet as rows, so every row can be found and cited on its own.

A CSV file and an Excel workbook are the same thing to a retriever: a header and
a list of records. What the retriever needs is each record as a chunk that says
what its values *are*, so a question about one employee's salary lands on that
employee's row and not on a block of numbers with the column names somewhere
above it.

Each data row becomes one `Document`:

    sheet: Sales            <- XLSX only
    month: March
    region: North
    revenue: 17747

with the row's place in the file in its metadata (`sheet`, `row`) so a citation
can say "Sales · row 14". `row` is the number a person sees in Excel or Google
Sheets -- the header is row 1, the first record row 2 -- because that is the
number they will go and look for.

Why this replaced langchain's CSV loader and the unstructured Excel loader. Both
were measured, on awkward files, before this was written:

  * `CSVLoader` does not detect a `;` delimiter, so a European-locale export
    arrived as one column named `id;name;city`. A duplicate header silently kept
    the later column and lost the earlier one's value; a short row embedded the
    word "None"; a byte-order mark stayed glued to the first column name.
  * `UnstructuredExcelLoader` flattened each sheet to one string with no row or
    column separators, so after chunking only the first chunk of a sheet had its
    header. On 400 rows, 8 of 15 salary lookups came back right, against 39 of 40
    on the same data as CSV. It also read hidden sheets, which someone who hides
    a sheet does not expect to be searchable.

No pandas and no unstructured: `csv` from the standard library, and `openpyxl` in
read-only mode.
"""

import csv
import io
import tempfile
from datetime import date, datetime, time
from pathlib import Path

from langchain_core.documents import Document

# A cell this short is part of what identifies a row (an id, a name, a
# department); a longer one is content. See `row_identity`.
IDENTITY_MAX_CELL = 80

# The identity block is repeated on the continuation chunks of a long row, so it
# is bounded: it must leave room for the text it introduces.
IDENTITY_MAX_CHARS = 300

# Python's default (128 KB) is smaller than a cell of pasted text can be.
csv.field_size_limit(10 * 1024 * 1024)


def row_identity(pairs: list[tuple[str, str]]) -> str:
    """The short cells of a row, as `column: value` lines, within the size cap.

    What a continuation chunk repeats so that it still says whose row it is. Every
    short cell is used, in column order, rather than guessing which column is the
    key: an id, a name and a city all identify a person, and which one a table
    leads with varies. Whole lines only -- a name cut in half identifies nobody.
    """
    lines: list[str] = []
    used = 0

    for column, value in pairs:
        if len(value) > IDENTITY_MAX_CELL:
            continue

        line = f"{column}: {value}"

        if used + len(line) + 1 > IDENTITY_MAX_CHARS:
            break

        lines.append(line)
        used += len(line) + 1

    return "\n".join(lines)


def continuation_prefix(metadata: dict, limit: int) -> str:
    """The header put in front of the second and later chunks of one row.

    `[sheet · ]row N (continued)`, then the row's identity. Shortened to *limit*
    characters, because the chunk it heads was already cut to the chunk size and
    the header must not double it.
    """
    where = f"row {metadata['row']} (continued)"

    if metadata.get("sheet"):
        where = f"{metadata['sheet']} · {where}"

    prefix = where
    identity = metadata.get("identity") or ""

    if identity:
        prefix = f"{where}\n{identity}"

    if len(prefix) > limit:
        prefix = prefix[: max(limit, len(where))].rstrip()

    return prefix + "\n"


def _header_names(raw: list[str]) -> list[str]:
    """Usable column names: no blanks, no duplicates, nothing lost.

    A blank header becomes `column_N` and a repeated one `name_2`, `name_3`, so
    two columns that share a name both keep their values.
    """
    seen: dict[str, int] = {}
    names: list[str] = []

    for index, cell in enumerate(raw, start=1):
        name = cell or f"column_{index}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        names.append(name if count == 1 else f"{name}_{count}")

    return names


def _pairs(columns: list[str], cells: list[str]) -> list[tuple[str, str]]:
    """Column/value pairs for one row, empty cells left out.

    A row longer than the header gets `column_N` names for the extra cells, and a
    shorter one simply has fewer pairs. Nothing is padded with a placeholder: an
    empty cell is the absence of a fact, and a literal "None" in the text is a
    fact a model will repeat back.
    """
    pairs = []

    for index, value in enumerate(cells):
        if not value:
            continue

        column = columns[index] if index < len(columns) else f"column_{index + 1}"
        pairs.append((column, value))

    return pairs


def _document(filename: str, sheet: str, row: int, pairs: list[tuple[str, str]]) -> Document:
    lines = [f"{column}: {value}" for column, value in pairs]

    if sheet:
        lines.insert(0, f"sheet: {sheet}")

    return Document(
        page_content="\n".join(lines),
        metadata={
            "source": filename,
            "sheet": sheet,
            "row": row,
            "identity": row_identity(pairs),
            "table": True,
        },
    )


def _decode(raw: bytes) -> str:
    """Text from CSV bytes: UTF-8 (a byte-order mark is dropped), else detected.

    Detection only on failure, because a file that decodes as UTF-8 is UTF-8 and
    the detector is slower and can be talked out of it by short Arabic text.
    """
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass

    try:
        from charset_normalizer import from_bytes

        best = from_bytes(raw).best()
        if best is not None:
            return str(best)
    except ImportError:
        pass

    return raw.decode("latin-1")


def _dialect(text: str) -> type[csv.Dialect]:
    """A dialect using the file's delimiter: the one that splits the header most.

    Counted on the first non-empty line rather than sniffed. `csv.Sniffer` needs
    several rows to decide and gets a lone header line wrong, and a header is the
    one line that always has every delimiter in it and no values that contain one.
    A comma when nothing else appears.
    """
    first = next((line for line in text.splitlines() if line.strip()), "")
    counts = {delimiter: first.count(delimiter) for delimiter in (",", ";", "\t", "|")}
    best = max(counts, key=lambda delimiter: (counts[delimiter], delimiter == ","))

    class Detected(csv.excel):
        delimiter = best if counts[best] else ","

    return Detected


def read_csv(path: Path) -> list[Document]:
    text = _decode(path.read_bytes())
    reader = csv.reader(io.StringIO(text, newline=""), _dialect(text))

    documents: list[Document] = []
    columns: list[str] | None = None

    for row_number, record in enumerate(reader, start=1):
        cells = [cell.strip() for cell in record]

        if not any(cells):
            continue

        if columns is None:
            columns = _header_names(cells)
            continue

        pairs = _pairs(columns, cells)

        if pairs:
            documents.append(_document(path.name, "", row_number, pairs))

    return documents


def _text(value) -> str:
    """One cell as the text a reader would see in the spreadsheet."""
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")

    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, time):
        return value.isoformat()

    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"

    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.15g}"

    return str(value).strip()


def read_xlsx(path: Path) -> list[Document]:
    from openpyxl import load_workbook

    # read_only streams the sheets instead of building them in memory, and
    # data_only reads the value a formula last calculated rather than its text.
    workbook = load_workbook(path, read_only=True, data_only=True)

    documents: list[Document] = []

    try:
        for sheet in workbook.worksheets:
            # A hidden sheet is hidden on purpose.
            if sheet.sheet_state != "visible":
                continue

            columns: list[str] | None = None

            for row_number, record in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = [_text(value) for value in record]

                if not any(cells):
                    continue

                if columns is None:
                    columns = _header_names(cells)
                    continue

                pairs = _pairs(columns, cells)

                if pairs:
                    documents.append(_document(path.name, sheet.title, row_number, pairs))
    finally:
        workbook.close()

    return documents


class TabularLoader:
    """A document loader for a CSV or XLSX file: `.load()` gives one Document per row."""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)

    def load(self) -> list[Document]:
        if self.path.suffix.lower() == ".xlsx":
            return read_xlsx(self.path)

        return read_csv(self.path)


def xlsx_preview_text(content: bytes) -> str:
    """A workbook's rows as plain text, for the source preview.

    The preview panel renders text; the bytes of an .xlsx are a zip archive and
    showed up as noise. Each row is headed by where it is -- the same
    `Sheet · row N` a citation names -- so a cited row can be found by eye.
    """
    with tempfile.NamedTemporaryFile(suffix=".xlsx") as handle:
        handle.write(content)
        handle.flush()
        documents = read_xlsx(Path(handle.name))

    blocks = []
    for document in documents:
        meta = document.metadata
        lines = [line for line in document.page_content.split("\n") if not line.startswith("sheet: ")]
        blocks.append(f"{meta['sheet']} · row {meta['row']}\n" + "\n".join(lines))

    return "\n\n".join(blocks)
